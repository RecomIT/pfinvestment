"""Phase 5: Excel exports.

Bond workbook (one per bond): every figure is a live Excel formula.
  Inputs     - bond and fund inputs (blue = typed values), IRR = XIRR(...) over the cash flows
  Schedule   - Calculation 1, effective interest; the last row is forced to land on face value
  Years / Quarters / Selected dates - Calculation 2 with a "position at each date" workings table
  TDS        - coupons received, TDS at the fund rate or the actual figure entered

Annexure-D workbook (one per fund and period): values from the app, totals as SUM formulas.

Only Excel-2007-era functions are used (XIRR, SUMIFS, COUNTIF, INDEX, MIN, IF, AND, OR) so the
files open the same in every Excel version and in LibreOffice.
"""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from typing import Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from engine import build_schedule
from engine.clock import today_local
from engine.fmt import dmy
from engine.models import DayBasis
from reports.annexure import Annexure
from storage.records import BondRecord, Fund

# ------------------------------------------------------------------ styles (house style of the client sheets)
FONT = "Arial"
F_BODY = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")          # typed values
F_LINK = Font(name=FONT, size=10, color="008000")           # pulled from another sheet
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=13, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
FILL_HEAD = PatternFill("solid", fgColor="1F4E79")
FILL_GROUP = PatternFill("solid", fgColor="DDEBF7")
FILL_TOTAL = PatternFill("solid", fgColor="F2F2F2")
FILL_CHECK = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
TOP2 = Border(top=Side(style="double", color="000000"))
WRAP = Alignment(wrap_text=True, vertical="bottom", horizontal="center")

NUM = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'          # same as the client's sheets
NUM2 = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
PCT2 = "0.00%"
PCT4 = "0.0000%"
DATE = "d-mmm-yy"
BASIS_TEXT = {DayBasis.EXCEL_182: "Excel 182", DayBasis.ACTUAL: "Actual"}


class Text(str):
    """A value typed by a user (name, note, heading). Always stored as plain text, never as a formula,
    so a name such as '=HYPERLINK(...)' or '=cmd|...' cannot run anything when the file is opened."""


def _set(ws, ref, value, font=F_BODY, fmt=None, fill=None, border=None, align=None):
    c = ws[ref]
    if isinstance(value, Text):
        c.value = str(value)
        c.data_type = "s"                     # force text even if it starts with = + - @
    else:
        c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = border
    if align:
        c.alignment = align
    return c


def _header_row(ws, row, labels, start_col=1, height=42):
    for i, lab in enumerate(labels):
        _set(ws, f"{get_column_letter(start_col + i)}{row}", lab, F_HEAD, fill=FILL_HEAD, border=BOX, align=WRAP)
    ws.row_dimensions[row].height = height


def _widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _print_setup(wb) -> None:
    """Landscape, A4, fit all columns on one page width (rows flow onto further pages)."""
    for ws in wb:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4


def _to_bytes(wb) -> bytes:
    _print_setup(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ====================================================================== bond workbook
def annexure_titles(fund_name: str, company_name: str) -> tuple[str, str]:
    """Heading lines: company on top when known, then the fund with the schedule name."""
    if company_name:
        return company_name, f"{fund_name} · Schedule of Bond Investment"
    return fund_name, "Schedule of Bond Investment"


def bond_workbook(rec: BondRecord, fund: Fund, selected: Optional[tuple[Optional[date], list[date]]] = None,
                  company_name: str = "") -> bytes:
    """Build the bond workbook. `selected` = (previous period end or None, cut-off dates) from the screen."""
    from engine.periods import fiscal_cutoffs

    b = rec.bond
    sched = build_schedule(b)
    n = len(sched.rows)
    wb = Workbook()

    # ------------------------------------------------------------------ Inputs
    ws = wb.active
    ws.title = "Inputs"
    _set(ws, "A1", Text(f"{b.isin}{' · ' + b.name if b.name else ''}"), F_TITLE)
    _set(ws, "A2", Text(f"{company_name + ' · ' if company_name else ''}{fund.name} ({fund.code}) · "
                        f"exported {dmy(today_local())}"), F_NOTE)
    _set(ws, "A3", "Blue = typed values you may change · black = formulas · green = from another sheet", F_NOTE)
    rows = [
        # label, name, value, format, font, note
        ("ISIN", None, Text(b.isin), None, F_INPUT, None),
        ("Name", None, Text(b.name), None, F_INPUT, None),
        ("Face value", "Face_Value", b.face_value, NUM, F_INPUT, None),
        ("Clean price paid", "Clean_Price", b.clean_price, NUM, F_INPUT, "Purchase price excluding accrued interest"),
        ("Accrued interest paid", "Accrued_Paid", b.accrued_interest_paid, NUM2, F_INPUT, "From the deal slip"),
        ("Settlement amount", None, "=Clean_Price+Accrued_Paid", NUM, F_BODY, None),
        ("Discount / (premium)", None, "=Face_Value-Clean_Price", NUM, F_BODY, None),
        ("Coupon rate", "Coupon_Rate", b.coupon_rate, PCT2, F_INPUT, None),
        ("Coupons per year", "Coupons_Per_Year", b.frequency, "0", F_INPUT, "Semi-annual only"),
        ("Coupon each period", "Coupon_Amount", "=Face_Value*Coupon_Rate/Coupons_Per_Year", NUM2, F_BODY, None),
        ("Issue date", "Issue_Date", b.issue_date, DATE, F_INPUT, None),
        ("Purchase date", "Purchase_Date", b.purchase_date, DATE, F_INPUT, None),
        ("Maturity date", "Maturity_Date", b.maturity_date, DATE, F_INPUT, None),
        ("Day basis at cut-off dates", "Day_Basis", BASIS_TEXT[b.day_basis], None, F_INPUT,
         "Excel 182: days since last coupon / 182. Actual: / actual days in the coupon period. "
         "The first (stub) period always uses actual stub days."),
        ("Excel half-year days", "Half_Year_Days", 182, "0", F_INPUT, "Fixed 182 days used by the Excel 182 basis"),
        ("IRR (XIRR of the cash flows)", "Bond_IRR",
         f"=XIRR(Schedule!$O$5:$O${n + 6},Schedule!$N$5:$N${n + 6})", PCT4, F_BODY,
         "Cash flows: -clean price on purchase; first coupon net of accrued interest paid; coupons; face at maturity"),
        ("IRR calculated by the app", None, sched.irr, PCT4, F_INPUT, "For reference only; not used by formulas"),
        ("Difference (Excel - app)", None, "=Bond_IRR-B21", "0.000000000%", F_BODY, None),
        ("Sold before maturity?", "Sold_Flag", "Yes" if b.sale else "No", None, F_INPUT, None),
        ("Sale date", "Sale_Date", b.sale.sale_date if b.sale else None, DATE, F_INPUT, None),
        ("Clean sale proceeds", "Sale_Clean", b.sale.clean_proceeds if b.sale else 0, NUM2, F_INPUT, None),
        ("Accrued interest received on sale", "Sale_Accrued", b.sale.accrued_interest_received if b.sale else 0, NUM2,
         F_INPUT, "Any difference from the calculated receivable goes to interest income"),
        ("Fund TDS rate on coupons", "TDS_Rate", fund.tds_rate, PCT2, F_INPUT, None),
    ]
    r0 = 5
    for i, (label, nm, value, fmt, font, note) in enumerate(rows):
        r = r0 + i
        _set(ws, f"A{r}", label, F_BOLD)
        c = _set(ws, f"B{r}", value, font, fmt)
        if nm:
            _name(wb, nm, f"Inputs!$B${r}")
        if note:
            _set(ws, f"C{r}", note, F_NOTE)
    # fixed positions used below (data validation) and by the difference formula
    assert ws["A18"].value == "Day basis at cut-off dates" and ws["A21"].value == "IRR calculated by the app"
    assert ws["A23"].value == "Sold before maturity?"
    basis_dv = DataValidation(type="list", formula1='"Excel 182,Actual"', allow_blank=False)
    sold_dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    ws.add_data_validation(basis_dv)
    ws.add_data_validation(sold_dv)
    basis_dv.add("B18")
    sold_dv.add("B23")
    if sched.warnings:
        r = r0 + len(rows) + 1
        _set(ws, f"A{r}", "Notes from the app", F_BOLD)
        for k, w in enumerate(sched.warnings, start=1):
            _set(ws, f"A{r + k}", Text(w), F_NOTE)
    _widths(ws, {"A": 34, "B": 20, "C": 80})
    ws.freeze_panes = "A5"

    # ------------------------------------------------------------------ Schedule
    sh = wb.create_sheet("Schedule")
    _set(sh, "A1", "Bond amortization schedule (effective interest method)", F_TITLE)
    _set(sh, "A2", "Interest = opening carrying value x ((1 + IRR) ^ (days / 365) - 1). "
                   "Amortization = interest - coupon received. The last row is set so the carrying value "
                   "ends exactly at face value.", F_NOTE)
    heads = ["#", "From", "Coupon date", "Days", "Beginning carrying value", "Interest (IRR)", "Coupon (gross)",
             "Accrued interest bought, recovered", "Coupon received (net)", "Amortization", "Cumulative amortization",
             "Ending carrying value"]
    _header_row(sh, 4, heads)
    _header_row(sh, 4, ["Cash flow date", "Cash flow (for XIRR)"], start_col=14)
    first, last = 5, 4 + n
    for i, row in enumerate(sched.rows):
        r = first + i
        is_last = i == n - 1
        _set(sh, f"A{r}", i + 1, F_BODY, "0")
        if i == 0:
            _set(sh, f"B{r}", "=Purchase_Date", F_LINK, DATE)
        else:
            _set(sh, f"B{r}", f"=C{r - 1}", F_BODY, DATE)
        _set(sh, f"C{r}", row.end, F_INPUT, DATE)
        _set(sh, f"D{r}", f"=C{r}-B{r}", F_BODY, "0")
        _set(sh, f"E{r}", "=Clean_Price" if i == 0 else f"=L{r - 1}", F_LINK if i == 0 else F_BODY, NUM)
        _set(sh, f"F{r}", f"=I{r}+J{r}" if is_last else f"=E{r}*((1+Bond_IRR)^(D{r}/365)-1)", F_BODY, NUM)
        _set(sh, f"G{r}", "=Coupon_Amount", F_LINK, NUM)
        _set(sh, f"H{r}", "=Accrued_Paid" if i == 0 else 0, F_LINK if i == 0 else F_BODY, NUM)
        _set(sh, f"I{r}", f"=G{r}-H{r}", F_BODY, NUM)
        _set(sh, f"J{r}", f"=Face_Value-E{r}" if is_last else f"=F{r}-I{r}", F_BODY, NUM)
        _set(sh, f"K{r}", f"=J{r}" if i == 0 else f"=K{r - 1}+J{r}", F_BODY, NUM)
        _set(sh, f"L{r}", f"=E{r}+J{r}", F_BODY, NUM)
    sh[f"F{last}"].comment = Comment("Last period: amortization = face value - opening, so the bond ends exactly "
                                     "at face value; interest = coupon + that amortization.", "Bond app")
    t = last + 1
    _set(sh, f"A{t}", "Total", F_BOLD, fill=FILL_TOTAL)
    for col in "FGHIJ":
        _set(sh, f"{col}{t}", f"=SUM({col}{first}:{col}{last})", F_BOLD, NUM, FILL_TOTAL)
    _set(sh, f"A{t + 2}", "Check: total amortization - discount/(premium) (should be 0)", F_BOLD)
    _set(sh, f"J{t + 2}", f"=J{t}-(Face_Value-Clean_Price)", F_BOLD, NUM2, FILL_CHECK)
    _set(sh, f"A{t + 3}", "Check: last ending carrying value - face value (should be 0)", F_BOLD)
    _set(sh, f"J{t + 3}", f"=L{last}-Face_Value", F_BOLD, NUM2, FILL_CHECK)
    # cash flows for XIRR: purchase, coupons (net), redemption
    _set(sh, "N5", "=Purchase_Date", F_LINK, DATE)
    _set(sh, "O5", "=-Clean_Price", F_LINK, NUM)
    for i in range(n):
        _set(sh, f"N{6 + i}", f"=C{first + i}", F_BODY, DATE)
        _set(sh, f"O{6 + i}", f"=I{first + i}", F_BODY, NUM)
    _set(sh, f"N{6 + n}", "=Maturity_Date", F_LINK, DATE)
    _set(sh, f"O{6 + n}", "=Face_Value", F_LINK, NUM)
    _widths(sh, {"A": 5, "B": 11, "C": 11, "D": 6, **{c: 15 for c in "EFGHIJKL"}, "M": 3, "N": 12, "O": 15})
    sh.freeze_panes = "B5"
    sh.print_title_rows = "4:4"
    rng = {k: f"Schedule!${col}${first}:${col}${last}" for k, col in
           {"start": "B", "end": "C", "days": "D", "interest": "F", "coupon": "G", "recovered": "H",
            "amort": "J"}.items()}

    # ------------------------------------------------------------------ period sheets
    close = b.sale.sale_date if b.sale else b.maturity_date
    before = b.purchase_date - timedelta(days=1)

    def through_close(ds: list[date]) -> list[date]:
        return [c for i, c in enumerate(ds) if i == 0 or ds[i - 1] < close]

    years = through_close(fiscal_cutoffs(b.purchase_date, close + timedelta(days=400),
                                         fund.year_end_month, fund.year_end_day, quarterly=False))
    quarters = through_close(fiscal_cutoffs(b.purchase_date, close + timedelta(days=400),
                                            fund.year_end_month, fund.year_end_day, quarterly=True))
    _period_sheet(wb, "Years", "Financial years", before, years, n, rng)
    _period_sheet(wb, "Quarters", "Quarters", before, quarters, n, rng)
    if selected and selected[1]:
        start, cuts = selected
        cuts = sorted(set(cuts))
        _period_sheet(wb, "Selected dates", "Dates chosen on screen", start or before, cuts, n, rng)

    # ------------------------------------------------------------------ TDS
    td = wb.create_sheet("TDS")
    _set(td, "A1", "Coupons received and TDS", F_TITLE)
    _set(td, "A2", "Actual TDS (blue) overrides the fund rate; leave blank to use the rate.", F_NOTE)
    _header_row(td, 4, ["Coupon date", "Gross coupon", "TDS rate", "TDS at fund rate", "Actual TDS entered",
                        "TDS used", "Net coupon received", "Note"])
    actual = {c.coupon_date: c for c in rec.coupon_actuals}
    received = [(i, r) for i, r in enumerate(sched.rows) if r.end <= close]
    for k, (i, row) in enumerate(received):
        r = 5 + k
        _set(td, f"A{r}", f"=Schedule!C{first + i}", F_LINK, DATE)
        _set(td, f"B{r}", f"=Schedule!G{first + i}", F_LINK, NUM)
        _set(td, f"C{r}", "=TDS_Rate", F_LINK, PCT2)
        _set(td, f"D{r}", f"=B{r}*C{r}", F_BODY, NUM)
        a = actual.get(row.end)
        _set(td, f"E{r}", a.tds if a else None, F_INPUT, NUM2)
        _set(td, f"F{r}", f'=IF(E{r}="",D{r},E{r})', F_BODY, NUM)
        _set(td, f"G{r}", f"=B{r}-F{r}", F_BODY, NUM)
        _set(td, f"H{r}", Text(a.note if a else ""), F_INPUT)
    if received:
        t = 5 + len(received)
        _set(td, f"A{t}", "Total", F_BOLD, fill=FILL_TOTAL)
        for col in "BDFG":
            _set(td, f"{col}{t}", f"=SUM({col}5:{col}{t - 1})", F_BOLD, NUM, FILL_TOTAL)
    _widths(td, {"A": 12, "B": 16, "C": 9, "D": 16, "E": 16, "F": 16, "G": 18, "H": 30})
    td.freeze_panes = "A5"

    return _to_bytes(wb)


def _period_sheet(wb, title: str, heading: str, start: date, cutoffs: list[date], n: int, rng: dict) -> None:
    """Calculation 2 as live formulas.

    Workings (lower table): the cumulative position at each date, exactly as engine.periods.state_at:
      partial coupon row k = (full periods ended) + 1; fraction = stub: actual stub days,
      Excel 182: MIN(days since last coupon / 182, 1), Actual: days / actual period days.
    Periods (upper table): differences between consecutive positions.
    """
    ws = wb.create_sheet(title)
    _set(ws, "A1", f"{heading}: interest income, amortization and carrying value", F_TITLE)
    _set(ws, "A2", "Each period runs from the day after the previous date to the cut-off date (inclusive). "
                   "Figures come from the 'Position at each date' workings below.", F_NOTE)
    m = len(cutoffs)
    dates = [start] + list(cutoffs)

    # ---- workings block position
    w_title = 4 + m + 4
    w_head = w_title + 1
    w0 = w_head + 1                       # first workings row (the opening date)
    wcols = ["Date", "Effective date (sale cap)", "Coupons ended", "Current period row", "Current period start",
             "Fraction of current period", "Cumulative interest", "Cumulative amortization",
             "Cumulative coupon (gross)", "Cumulative bought interest recovered", "Receivable before sale",
             "Sold by this date (1/0)", "Invested", "Matured (1/0)", "Carrying before closing", "Encashed principal",
             "Carrying value", "Sale interest adjustment", "Interest income (cumulative)",
             "Accrued interest on sale", "Gain / (loss) on sale", "Interest receivable"]
    _set(ws, f"A{w_title}", "Workings: position at each date (cumulative from purchase)", F_BOLD)
    _header_row(ws, w_head, wcols, height=54)
    S, E, D, I_, C_, R_, A_ = (rng[k] for k in ("start", "end", "days", "interest", "coupon", "recovered", "amort"))
    for j, d in enumerate(dates):
        r = w0 + j
        f = {
            "A": d,
            "B": f'=IF(AND(Sold_Flag="Yes",A{r}>=Sale_Date),Sale_Date,A{r})',
            "C": f'=COUNTIF({E},"<="&B{r})',
            "D": f"=C{r}+1",
            "E": f'=IF(D{r}<={n},INDEX({S},D{r}),"")',
            "F": (f'=IF(D{r}>{n},0,IF(B{r}<=E{r},0,IF(D{r}=1,(B{r}-E{r})/INDEX({D},D{r}),'
                  f'IF(Day_Basis="Excel 182",MIN((B{r}-E{r})/Half_Year_Days,1),(B{r}-E{r})/INDEX({D},D{r})))))'),
            "G": f'=SUMIFS({I_},{E},"<="&B{r})+IF(D{r}<={n},F{r}*INDEX({I_},D{r}),0)',
            "H": f'=SUMIFS({A_},{E},"<="&B{r})+IF(D{r}<={n},F{r}*INDEX({A_},D{r}),0)',
            "I": f'=SUMIFS({C_},{E},"<="&B{r})',
            "J": f'=SUMIFS({R_},{E},"<="&B{r})',
            "K": f"=G{r}-H{r}-(I{r}-J{r})",
            "L": f'=IF(AND(Sold_Flag="Yes",A{r}>=Sale_Date),1,0)',
            "M": f"=IF(A{r}>=Purchase_Date,Clean_Price,0)",
            "N": f"=IF(A{r}>=Maturity_Date,1,0)",
            "O": f"=IF(M{r}=0,0,Clean_Price+H{r})",
            "P": f"=IF(OR(L{r}=1,N{r}=1),O{r},0)",
            "Q": f"=IF(OR(L{r}=1,N{r}=1),0,O{r})",
            "R": f"=IF(L{r}=1,Sale_Accrued-K{r},0)",
            "S": f"=G{r}+R{r}",
            "T": f"=IF(L{r}=1,Sale_Accrued,0)",
            "U": f"=IF(L{r}=1,Sale_Clean-O{r},0)",
            "V": f"=IF(L{r}=1,0,K{r})",
        }
        for col, v in f.items():
            fmt = DATE if col in "ABE" else ("0" if col in "CDLN" else ("0.000000" if col == "F" else NUM))
            _set(ws, f"{col}{r}", v, F_INPUT if col == "A" else F_BODY, fmt)

    # ---- periods table
    heads = ["Period end", "Opening carrying value", "Investment", "Amortization", "Encashed",
             "Closing carrying value", "Opening interest receivable", "Interest income", "Coupon received (gross)",
             "Holding period interest", "Accrued interest on sale", "Gain / (loss) on sale",
             "Closing interest receivable", "Check: carrying (0)", "Check: receivable (0)"]
    _set(ws, "A3", f"Periods (opening date {dmy(start)})", F_BOLD)
    _header_row(ws, 4, heads, height=54)
    for j in range(m):
        r, cur, prev = 5 + j, w0 + j + 1, w0 + j
        row = {
            "A": f"=A{cur}",
            "B": f"=Q{prev}",
            "C": f"=M{cur}-M{prev}",
            "D": f"=H{cur}-H{prev}",
            "E": f"=P{cur}-P{prev}",
            "F": f"=Q{cur}",
            "G": f"=V{prev}",
            "H": f"=S{cur}-S{prev}",
            "I": f"=I{cur}-I{prev}",
            "J": f"=J{cur}-J{prev}",
            "K": f"=T{cur}-T{prev}",
            "L": f"=U{cur}-U{prev}",
            "M": f"=V{cur}",
            "N": f"=B{r}+C{r}+D{r}-E{r}-F{r}",
            "O": f"=G{r}+H{r}-D{r}-(I{r}-J{r})-K{r}-M{r}",
        }
        for col, v in row.items():
            _set(ws, f"{col}{r}", v, F_BODY, DATE if col == "A" else (NUM2 if col in "NO" else NUM),
                 FILL_CHECK if col in "NO" else None)
    if m:
        t = 5 + m
        _set(ws, f"A{t}", "Total", F_BOLD, fill=FILL_TOTAL)
        for col in "CDEHIJKL":
            _set(ws, f"{col}{t}", f"=SUM({col}5:{col}{t - 1})", F_BOLD, NUM, FILL_TOTAL)
    _widths(ws, {"A": 12, **{get_column_letter(k): 15 for k in range(2, 23)}})
    ws.freeze_panes = "B5"


# ====================================================================== Annexure-D workbook
ANNEX_COLS = [
    # letter, header, attribute, format, total?
    ("A", "Serial Number", "serial", "0", False), ("B", "ISIN", "isin", None, False),
    ("C", "Issue Date", "issue_date", DATE, False), ("D", "Purchase Date", "purchase_date", DATE, False),
    ("E", "Maturity Date", "maturity_date", DATE, False), ("F", "Nature", "nature", None, False),
    ("G", "Coupon Rate", "coupon_rate", PCT2, False), ("H", "IRR", "irr", PCT4, False),
    ("I", "Face Value", "face_value", NUM, True), ("J", "Purchase/ Settlement price", "settlement_amount", NUM, True),
    ("K", "Opening investment", "opening_investment", NUM, True),
    ("L", "Investment During the Period", "investment", NUM, True),
    ("M", "Amortization Discount/premium", "amortization", NUM, True), ("N", "Encashed", "encashed", NUM, True),
    ("O", "Closing Balance", "closing_balance", NUM, True), ("P", "Opening interest", "opening_interest", NUM, True),
    ("Q", "Interest income during the period", "interest_income", NUM, True),
    ("R", "Holding Period Interest", "holding_period_interest", NUM, True),
    ("S", "Gross interest receipt", "gross_interest_receipt", NUM, True),
    ("T", "Interest Receipt", "interest_receipt", NUM, True), ("U", "TDS", "tds", NUM, True),
    ("V", "Encashed", "encashed_interest", NUM, True),
    ("W", "Closing interest receivable", "closing_interest_receivable", NUM, True),
    ("X", "Gain/(Loss) on sale", "gain_loss", NUM, True),
]


def annexure_workbook(a: Annexure) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Annexure-D"
    last_col = ANNEX_COLS[-1][0]
    line1, line2 = annexure_titles(a.fund.name, a.company_name)
    for r, text, font in ((1, Text(line1), F_TITLE), (2, Text(line2), F_BOLD),
                          (3, f"For the period {dmy(a.period_start + timedelta(days=1))} to {dmy(a.period_end)}",
                           F_BOLD)):
        ws.merge_cells(f"A{r}:{last_col}{r}")
        _set(ws, f"A{r}", text, font, align=Alignment(horizontal="center"))
    for rng_, label in (("A5:J5", "Particulars"), ("K5:O5", "Principal"), ("P5:W5", "Interest"), ("X5:X5", "Sale")):
        if ":" in rng_ and rng_.split(":")[0] != rng_.split(":")[1]:
            ws.merge_cells(rng_)
        _set(ws, rng_.split(":")[0], label, F_BOLD, fill=FILL_GROUP, align=Alignment(horizontal="center"))
    _header_row(ws, 6, [c[1] for c in ANNEX_COLS], height=54)
    for c in ANNEX_COLS:                                  # column letters, like the sheet's own references
        _set(ws, f"{c[0]}7", c[0], F_NOTE, align=Alignment(horizontal="center"))
    first = 8
    for i, row in enumerate(a.rows):
        r = first + i
        for letter, _, attr, fmt, _t in ANNEX_COLS:
            v = getattr(row, attr)
            _set(ws, f"{letter}{r}", Text(v) if isinstance(v, str) else v, F_BODY, fmt)
    last = first + len(a.rows) - 1
    t = last + 1 if a.rows else first
    _set(ws, f"A{t}", "Total", F_BOLD, fill=FILL_TOTAL)
    for letter, _, _, fmt, tot in ANNEX_COLS:
        if tot:
            _set(ws, f"{letter}{t}", f"=SUM({letter}{first}:{letter}{last})" if a.rows else 0, F_BOLD, fmt,
                 FILL_TOTAL, TOP2)
    # arithmetic checks, same as the sheet's own formulas (should all be 0)
    c = t + 2
    _set(ws, f"A{c}", "Checks (should be 0): O = K + L + M - N;  W = P + Q - M - S + R - V;  S = T + U", F_NOTE)
    if a.rows:
        _set(ws, f"O{c}", f"=K{t}+L{t}+M{t}-N{t}-O{t}", F_BOLD, NUM2, FILL_CHECK)
        _set(ws, f"W{c}", f"=P{t}+Q{t}-M{t}-S{t}+R{t}-V{t}-W{t}", F_BOLD, NUM2, FILL_CHECK)
        _set(ws, f"S{c}", f"=T{t}+U{t}-S{t}", F_BOLD, NUM2, FILL_CHECK)
    if a.warnings:
        _set(ws, f"A{c + 2}", "Notes", F_BOLD)
        for k, w in enumerate(a.warnings, start=1):
            _set(ws, f"A{c + 2 + k}", Text(w), F_NOTE)
    _set(ws, f"A{c + 1}", f"Amounts in Taka. Values calculated by the bond app on {dmy(today_local())}.", F_NOTE)
    _widths(ws, {"A": 8, "B": 15, "C": 11, "D": 11, "E": 11, "F": 10, "G": 9, "H": 10,
                 **{get_column_letter(k): 15 for k in range(9, 25)}})
    ws.freeze_panes = "C8"
    ws.print_title_rows = "5:7"
    return _to_bytes(wb)
